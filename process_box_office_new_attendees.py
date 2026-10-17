import os
import sys
import io
import csv
import shutil
from datetime import datetime
from PIL import Image, ImageDraw, ImageFont
import openpyxl
import barcode
from barcode.writer import ImageWriter
import fitz  # PyMuPDF
from tinydb import TinyDB, Query

# Ensure UTF-8 output encoding on Windows consoles
if sys.stdout.encoding != 'utf-8':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DANDIYAA_DIR = os.path.join(BASE_DIR, 'dandiyaa')
DESIGNS_DIR = os.path.join(DANDIYAA_DIR, 'Pass_designs')
SHEET_PATH = os.path.join(DANDIYAA_DIR, 'Sheet.xlsx')

DANDIYA_TEMPLATE = os.path.join(DESIGNS_DIR, 'BOX OFFICE PASS DANDIYA.jpg.jpeg')
CARNIVAL_TEMPLATE = os.path.join(BASE_DIR, 'BOX_OFFICE.jpeg')

NEW_ATTENDEES = [
    {'idx': 1,  'name': 'YASH SACHIN BONDE',            'id': '32251'},
    {'idx': 2,  'name': 'RAUL RAJ SHANKARSING',        'id': '32211'},
    {'idx': 3,  'name': 'RANARANJEET RAVINDRA KULKARNI','id': '32208'},
    {'idx': 4,  'name': 'RITESH VIJAY PATIL',           'id': '32213'},
    {'idx': 5,  'name': 'CHANDANSHIV YASH VIJAY',       'id': '22118'},
    {'idx': 6,  'name': 'PRATIK TARATE',               'id': '11239'},
    {'idx': 7,  'name': 'YUGANDHAR TUPE',               'id': '22264'},
    {'idx': 8,  'name': 'VIRAJ AMOLIK',                'id': '24104'},
    {'idx': 9,  'name': 'AADITYA SALVE',               'id': '12104'},
    {'idx': 10, 'name': 'PRATHMESH SHINGANE',           'id': '12322'},
    {'idx': 11, 'name': 'ANISHA PRIYA',                'id': '13109'},
    {'idx': 12, 'name': 'SAYALI AHIRE',                'id': '13107'},
    {'idx': 13, 'name': 'SAKSHI BHANUSHALI',           'id': '13118'},
    {'idx': 14, 'name': 'SUDARSHAN GOSAVI',            'id': '13171'},
    {'idx': 15, 'name': 'HEMANT PATIL',                'id': '13301'},
    {'idx': 16, 'name': 'SHRUTI PENDHARE',             'id': '13310'},
    {'idx': 17, 'name': 'SHWET THARVAL',               'id': '316508907510'},
]

def get_font_path():
    candidates = [
        r'C:\Windows\Fonts\arialbd.ttf',
        r'C:\Windows\Fonts\segoeuib.ttf',
        r'C:\Windows\Fonts\calibrib.ttf'
    ]
    for p in candidates:
        if os.path.exists(p):
            return p
    return candidates[0]

def render_dandiya_pass(template_img, name, student_id, font_path, code128_class, writer):
    """
    Renders high-res Dandiya Box Office pass (3000x900).
    Cleanly centers Name, Roll/Student ID, and Code-128 Barcode.
    """
    pass_img = template_img.copy()
    draw = ImageDraw.Draw(pass_img)
    center_x = 2419
    max_name_width = 820

    # Dynamic font sizing for student name
    font_size_name = 100
    font_name = ImageFont.truetype(font_path, font_size_name)
    bb = draw.textbbox((0, 0), name, font=font_name, stroke_width=2)
    while (bb[2] - bb[0]) > max_name_width and font_size_name > 30:
        font_size_name -= 3
        font_name = ImageFont.truetype(font_path, font_size_name)
        bb = draw.textbbox((0, 0), name, font=font_name, stroke_width=2)

    # Dynamic font sizing for student ID
    font_size_id = 120
    font_id = ImageFont.truetype(font_path, font_size_id)
    bb_i = draw.textbbox((0, 0), student_id, font=font_id, stroke_width=2)
    while (bb_i[2] - bb_i[0]) > max_name_width and font_size_id > 40:
        font_size_id -= 4
        font_id = ImageFont.truetype(font_path, font_size_id)
        bb_i = draw.textbbox((0, 0), student_id, font=font_id, stroke_width=2)

    # Coordinates
    bb_n = draw.textbbox((0, 0), name, font=font_name, stroke_width=2)
    wn, hn = bb_n[2] - bb_n[0], bb_n[3] - bb_n[1]
    xn = center_x - wn // 2
    yn = 405 - hn // 2

    wi, hi = bb_i[2] - bb_i[0], bb_i[3] - bb_i[1]
    xi = center_x - wi // 2
    yi = 560 - hi // 2

    # Draw Name & ID with crisp white stroke
    draw.text((xn, yn), name, fill=(255, 255, 255), stroke_width=2, stroke_fill=(255, 255, 255), font=font_name)
    draw.text((xi, yi), student_id, fill=(255, 255, 255), stroke_width=2, stroke_fill=(255, 255, 255), font=font_id)

    # Barcode (Code 128)
    bc_opts = {
        'module_width': 0.40,
        'module_height': 14.0,
        'quiet_zone': 2.0,
        'font_size': 0,
        'text_distance': 0,
        'write_text': False
    }
    bc = code128_class(student_id, writer=writer)
    fp = io.BytesIO()
    bc.write(fp, options=bc_opts)
    fp.seek(0)
    bc_img = Image.open(fp)

    box_w = 580
    box_h = 135
    box_center_x = 2419
    box_center_y = 754

    bw, bh = bc_img.size
    scale = min(box_w / bw, box_h / bh)
    new_bw = max(1, int(bw * scale))
    new_bh = max(1, int(bh * scale))
    bc_resized = bc_img.resize((new_bw, new_bh), Image.Resampling.LANCZOS)

    paste_x = box_center_x - new_bw // 2
    paste_y = box_center_y - new_bh // 2
    pass_img.paste(bc_resized, (paste_x, paste_y))

    return pass_img

def render_carnival_pass(template_img, name, student_id, font_path, code128_class, writer):
    """
    Renders high-res Carnival Box Office pass (3000x900).
    Cleanly centers Name, Roll/Student ID, and Code-128 Barcode with elegant styling.
    """
    pass_img = template_img.copy()
    draw = ImageDraw.Draw(pass_img)
    center_x = 2460
    max_name_width = 850

    font_size_name = 85
    font_name = ImageFont.truetype(font_path, font_size_name)
    bb = draw.textbbox((0, 0), name, font=font_name)
    while (bb[2] - bb[0]) > max_name_width and font_size_name > 30:
        font_size_name -= 3
        font_name = ImageFont.truetype(font_path, font_size_name)
        bb = draw.textbbox((0, 0), name, font=font_name)

    font_id = ImageFont.truetype(font_path, 68)

    bb_n = draw.textbbox((0, 0), name, font=font_name)
    wn, hn = bb_n[2] - bb_n[0], bb_n[3] - bb_n[1]
    xn = center_x - wn // 2
    yn = 420 - hn // 2

    serial_text = f'Roll No. {student_id}'
    bb_i = draw.textbbox((0, 0), serial_text, font=font_id)
    wi, hi = bb_i[2] - bb_i[0], bb_i[3] - bb_i[1]
    xi = center_x - wi // 2
    yi = 545 - hi // 2

    draw.text((xn, yn), name, fill=(20, 15, 30), font=font_name)
    draw.text((xi, yi), serial_text, fill=(160, 95, 10), font=font_id)

    bc = code128_class(student_id, writer=writer)
    fp = io.BytesIO()
    bc.write(fp, options={
        'module_width': 0.40,
        'module_height': 14.0,
        'quiet_zone': 2.0,
        'font_size': 0,
        'text_distance': 0,
        'write_text': False
    })
    fp.seek(0)
    bc_img = Image.open(fp).convert('RGB')

    box_w = 760
    box_h = 150
    box_x1 = center_x - box_w // 2
    box_y1 = 665
    box_x2 = box_x1 + box_w
    box_y2 = box_y1 + box_h
    draw.rectangle([box_x1, box_y1, box_x2, box_y2], fill=(255, 255, 255), outline=(190, 150, 80), width=3)

    bw, bh = bc_img.size
    scale = min(660 / bw, 120 / bh)
    new_bw = max(1, int(bw * scale))
    new_bh = max(1, int(bh * scale))
    bc_resized = bc_img.resize((new_bw, new_bh), Image.Resampling.LANCZOS)

    bc_x = box_x1 + (box_w - new_bw) // 2
    bc_y = box_y1 + (box_h - new_bh) // 2
    pass_img.paste(bc_resized, (bc_x, bc_y))

    return pass_img

def build_pdf_documents(pass_images, output_a4_path, output_strip_path):
    """
    Builds both A4 Printable PDF (3 passes/page) and Multi-Page Strip PDF.
    """
    # 1. A4 Printable
    doc_a4 = fitz.open()
    a4_w, a4_h = 595.28, 841.89  # points
    pass_w = 555.0
    pass_h = pass_w * (900.0 / 3000.0)  # 166.5 pt
    margin_x = (a4_w - pass_w) / 2.0
    spacing_y = 60.0
    start_y = (a4_h - (3 * pass_h + 2 * spacing_y)) / 2.0

    current_page = None
    for i, img_bytes in enumerate(pass_images):
        slot_idx = i % 3
        if slot_idx == 0:
            current_page = doc_a4.new_page(width=a4_w, height=a4_h)

        slot_y = start_y + slot_idx * (pass_h + spacing_y)
        rect_slot = fitz.Rect(margin_x, slot_y, margin_x + pass_w, slot_y + pass_h)
        current_page.insert_image(rect_slot, stream=img_bytes)

        border_rect = fitz.Rect(margin_x - 0.5, slot_y - 0.5, margin_x + pass_w + 0.5, slot_y + pass_h + 0.5)
        current_page.draw_rect(border_rect, color=(0.8, 0.8, 0.8), width=0.5)

    doc_a4.save(output_a4_path, deflate=True)
    doc_a4.close()

    # 2. Multi-Page Strip PDF (720 x 216 pt)
    doc_strip = fitz.open()
    strip_w, strip_h = 720, 216
    rect_single = fitz.Rect(0, 0, strip_w, strip_h)

    for img_bytes in pass_images:
        p = doc_strip.new_page(width=strip_w, height=strip_h)
        p.insert_image(rect_single, stream=img_bytes)

    doc_strip.save(output_strip_path, deflate=True)
    doc_strip.close()

def append_to_sheet(sheet_path, attendees):
    """
    Appends new attendees to dandiyaa/Sheet.xlsx under the BOX OFFICE section.
    """
    if not os.path.exists(sheet_path):
        print(f"[Sheet] Warning: {sheet_path} does not exist, skipping excel append.")
        return

    wb = openpyxl.load_workbook(sheet_path)
    ws = wb['Sheet1']

    # Find the last row with data
    last_row = 0
    for r in range(ws.max_row, 0, -1):
        if any(ws.cell(r, c).value is not None for c in range(1, 5)):
            last_row = r
            break

    existing_ids = set()
    for r in range(1, last_row + 1):
        val = ws.cell(r, 4).value
        if val is not None:
            existing_ids.add(str(val).strip().upper())

    added_count = 0
    current_row = last_row + 1

    for a in attendees:
        sid_str = str(a['id']).strip().upper()
        # Col A: None (under BOX OFFICE header)
        # Col B: Attendee Name
        # Col C: Division / Group ('BOX OFFICE')
        # Col D: Student / Roll ID
        ws.cell(row=current_row, column=1, value=None)
        ws.cell(row=current_row, column=2, value=a['name'])
        ws.cell(row=current_row, column=3, value='BOX OFFICE')
        ws.cell(row=current_row, column=4, value=a['id'])
        current_row += 1
        added_count += 1

    wb.save(sheet_path)
    print(f"[Sheet] Appended {added_count} attendees to {sheet_path} (rows {last_row + 1} to {current_row - 1}).")

def sync_to_database(db_path, attendees):
    """
    Syncs attendees into TinyDB students table.
    Ensures barcode, name, college, division, roll_no, status='PENDING'.
    Also updates any shorter matching names in other colleges (e.g. 'RAJ' -> 'RAUL RAJ SHANKARSING').
    """
    if not os.path.exists(db_path):
        print(f"[DB] {db_path} does not exist, skipping.")
        return 0, 0

    db = TinyDB(db_path)
    students_table = db.table('students')
    all_students = students_table.all()
    q = Query()

    # Index by (college, barcode)
    existing_by_col_bc = {}
    for s in all_students:
        c = (s.get('college') or '').strip().upper()
        bc = (s.get('barcode') or '').strip().upper()
        if c and bc:
            existing_by_col_bc[(c, bc)] = s

    inserted = 0
    updated = 0

    for a in attendees:
        bc = str(a['id']).strip().upper()
        name = a['name'].strip().upper()
        key = ('BOX OFFICE', bc)

        if key in existing_by_col_bc:
            # Update name if more detailed
            cur_record = existing_by_col_bc[key]
            if len(name) > len(cur_record.get('name', '')):
                students_table.update(
                    {'name': name, 'roll_no': bc},
                    (q.college == 'BOX OFFICE') & (q.barcode == bc)
                )
                updated += 1
        else:
            record = {
                'barcode': bc,
                'name': name,
                'college': 'BOX OFFICE',
                'division': 'BOX OFFICE',
                'roll_no': bc,
                'status': 'PENDING',
                'entry_timestamp': None,
                'entry_date': None,
                'entry_time': None,
                'volunteer': None
            }
            students_table.insert(record)
            existing_by_col_bc[key] = record
            inserted += 1

        # Also update other colleges if they had truncated names for the same student roll number
        for other_col in ['DYPCOEI', 'DYPIMED', 'DYPTC', 'DYPSOA']:
            other_key = (other_col, bc)
            if other_key in existing_by_col_bc:
                other_rec = existing_by_col_bc[other_key]
                if len(name) > len(other_rec.get('name', '')):
                    students_table.update(
                        {'name': name},
                        (q.college == other_col) & (q.barcode == bc)
                    )
                    updated += 1

    db.close()
    db_name = os.path.basename(os.path.dirname(db_path)) or 'root'
    print(f"[DB] {db_name}/{os.path.basename(db_path)}: Inserted {inserted} new records, updated {updated} names.")
    return inserted, updated

def main():
    print("=" * 70)
    print("BOX OFFICE PASS GENERATION & DATABASE INTEGRATION")
    print(f"Total new attendees to process: {len(NEW_ATTENDEES)}")
    print("=" * 70)

    font_path = get_font_path()
    c128_class = barcode.get_barcode_class('code128')
    writer = ImageWriter()

    # 1. Update dandiyaa/Sheet.xlsx
    append_to_sheet(SHEET_PATH, NEW_ATTENDEES)

    # 2. Sync to TinyDB databases
    sync_to_database(os.path.join(DANDIYAA_DIR, 'database.json'), NEW_ATTENDEES)
    sync_to_database(os.path.join(BASE_DIR, 'database.json'), NEW_ATTENDEES)

    # 3. Export CSV files
    csv_root = os.path.join(BASE_DIR, 'BOX_OFFICE_New_17_Attendees.csv')
    csv_dandiya = os.path.join(DANDIYAA_DIR, 'BOX_OFFICE_New_17_Attendees.csv')
    for cp in [csv_root, csv_dandiya]:
        with open(cp, 'w', newline='', encoding='utf-8') as f:
            cw = csv.writer(f)
            cw.writerow(['SrNo', 'Attendee Name', 'Category/College', 'Division', 'Roll No / Barcode'])
            for a in NEW_ATTENDEES:
                cw.writerow([a['idx'], a['name'], 'BOX OFFICE', 'BOX OFFICE', a['id']])
    print(f"[CSV] Created attendee CSV at: {csv_root} and {csv_dandiya}")

    # 4. Generate Dandiya Box Office Passes
    print("\n--- Generating Dandiya Box Office Passes ---")
    t_dandiya = Image.open(DANDIYA_TEMPLATE).convert('RGB')
    dandiya_images = []
    
    # Directory for individual images
    img_dir_dandiya = os.path.join(BASE_DIR, 'passes', 'dandiyaa_box_office_new_17')
    os.makedirs(img_dir_dandiya, exist_ok=True)

    for a in NEW_ATTENDEES:
        p_img = render_dandiya_pass(t_dandiya, a['name'], a['id'], font_path, c128_class, writer)
        
        # Save individual pass PNG
        safe_name = a['name'].replace(' ', '_')
        fn = f"{a['idx']:02d}_{safe_name}_{a['id']}.png"
        p_img.save(os.path.join(img_dir_dandiya, fn))

        buf = io.BytesIO()
        p_img.save(buf, format='JPEG', quality=92)
        dandiya_images.append(buf.getvalue())

    dandiya_a4_pdf = os.path.join(BASE_DIR, 'BOX_OFFICE_New_17_Passes_A4_Printable.pdf')
    dandiya_strip_pdf = os.path.join(BASE_DIR, 'BOX_OFFICE_New_17_Strip_Passes.pdf')
    build_pdf_documents(dandiya_images, dandiya_a4_pdf, dandiya_strip_pdf)

    # Also copy to dandiyaa/ directory
    shutil.copy2(dandiya_a4_pdf, os.path.join(DANDIYAA_DIR, 'BOX_OFFICE_New_17_Passes_A4_Printable.pdf'))
    shutil.copy2(dandiya_strip_pdf, os.path.join(DANDIYAA_DIR, 'BOX_OFFICE_New_17_Strip_Passes.pdf'))

    print(f"[Dandiya PDF] Created: {dandiya_a4_pdf} ({len(NEW_ATTENDEES)} passes on {(len(NEW_ATTENDEES)+2)//3} A4 pages)")
    print(f"[Dandiya PDF] Created: {dandiya_strip_pdf} ({len(NEW_ATTENDEES)} strip passes)")
    print(f"[Dandiya Images] Saved 17 individual pass images to: {img_dir_dandiya}")

    # 5. Generate Carnival Box Office Passes
    print("\n--- Generating Carnival Box Office Passes ---")
    t_carnival = Image.open(CARNIVAL_TEMPLATE).convert('RGB')
    carnival_images = []
    
    img_dir_carnival = os.path.join(BASE_DIR, 'passes', 'carnival_box_office_new_17')
    os.makedirs(img_dir_carnival, exist_ok=True)

    for a in NEW_ATTENDEES:
        p_img = render_carnival_pass(t_carnival, a['name'], a['id'], font_path, c128_class, writer)
        safe_name = a['name'].replace(' ', '_')
        fn = f"{a['idx']:02d}_{safe_name}_{a['id']}.png"
        p_img.save(os.path.join(img_dir_carnival, fn))

        buf = io.BytesIO()
        p_img.save(buf, format='JPEG', quality=92)
        carnival_images.append(buf.getvalue())

    carnival_a4_pdf = os.path.join(BASE_DIR, 'Carnival_BOX_OFFICE_New_17_Passes_A4_Printable.pdf')
    carnival_strip_pdf = os.path.join(BASE_DIR, 'Carnival_BOX_OFFICE_New_17_Strip_Passes.pdf')
    build_pdf_documents(carnival_images, carnival_a4_pdf, carnival_strip_pdf)

    print(f"[Carnival PDF] Created: {carnival_a4_pdf} ({len(NEW_ATTENDEES)} passes on {(len(NEW_ATTENDEES)+2)//3} A4 pages)")
    print(f"[Carnival PDF] Created: {carnival_strip_pdf} ({len(NEW_ATTENDEES)} strip passes)")
    print(f"[Carnival Images] Saved 17 individual pass images to: {img_dir_carnival}")

    # 6. Regenerate Complete Box Office Dandiya Passes (Including previous + new = 225 passes)
    print("\n--- Updating Complete Dandiya Box Office Master Passes ---")
    import generate_dandiya_passes
    recs_by_col = generate_dandiya_passes.load_sheet_records()
    box_office_all = recs_by_col.get('BOX OFFICE', [])
    print(f"[Master] Total BOX OFFICE records loaded from updated Sheet: {len(box_office_all)}")
    
    res = generate_dandiya_passes.generate_college_passes(
        'BOX OFFICE',
        box_office_all,
        font_path,
        c128_class,
        writer
    )
    print(f"[Master] Successfully updated complete Dandiya Box Office passes: {res['a4_pdf']} ({res['a4_pages']} pages, {res['count']} passes)")

    print("\n" + "=" * 70)
    print("ALL 17 BOX OFFICE PASSES & DATABASE RECORDS GENERATED SUCCESSFULLY!")
    print("=" * 70)

if __name__ == '__main__':
    main()
