import os
import sys
import io
import shutil
import openpyxl

# Set console output encoding for Windows compatibility
if sys.stdout.encoding != 'utf-8':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass
from PIL import Image, ImageDraw, ImageFont
import barcode
from barcode.writer import ImageWriter
import fitz  # PyMuPDF
from tinydb import TinyDB, Query

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DANDIYAA_DIR = os.path.join(BASE_DIR, 'dandiyaa')
DESIGNS_DIR = os.path.join(DANDIYAA_DIR, 'Pass_designs')
PASSES_COPY_DIR = os.path.join(BASE_DIR, 'passes', 'dandiyaa')

SHEET_PATH = os.path.join(DANDIYAA_DIR, 'Sheet.xlsx')

TEMPLATES = {
    'BOX OFFICE': os.path.join(DESIGNS_DIR, 'BOX OFFICE PASS DANDIYA.jpg.jpeg'),
    'DYPCOEI': os.path.join(DESIGNS_DIR, 'DYPCOEI PASS DANDIYA.jpg.jpeg'),
    'DYPIMED': os.path.join(DESIGNS_DIR, 'DYPIMED PASS DANDIYA.jpg.jpeg'),
    'DYPSOA': os.path.join(DESIGNS_DIR, 'DYPSOA PASS DANDIYA.jpg.jpeg'),
    'DYPTC': os.path.join(DESIGNS_DIR, 'DYPTC PASS DANDIYA.jpg.jpeg'),
}

def get_font_path():
    candidates = [
        r'C:\Windows\Fonts\arialbd.ttf',
        r'C:\Windows\Fonts\segoeuib.ttf',
        r'C:\Windows\Fonts\calibrib.ttf'
    ]
    for p in candidates:
        if os.path.exists(p):
            return p
    return candidates[0]

def load_sheet_records():
    """
    Parses dandiyaa/Sheet.xlsx and groups student records into the 5 colleges:
    DYPIMED, DYPCOEI, DYPTC, DYPSOA, BOX OFFICE.
    """
    wb = openpyxl.load_workbook(SHEET_PATH, data_only=True)
    ws = wb['Sheet1']

    records_by_college = {
        'DYPIMED': [],
        'DYPCOEI': [],
        'DYPTC': [],
        'DYPSOA': [],
        'BOX OFFICE': []
    }
    current_col = None

    for r in range(1, ws.max_row + 1):
        vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        c0 = str(vals[0]).strip() if vals[0] is not None else ''

        # Detect section header
        for known in ['DYPIMED', 'DYPCOEI', 'DYPTC', 'DYPSOA', 'BOX OFFICE']:
            if c0 == known or (known in c0 and len([x for x in vals if x is not None]) <= 2):
                current_col = known
                break

        name = str(vals[1]).strip() if vals[1] is not None else ''
        div = str(vals[2]).strip() if vals[2] is not None else ''
        sid = str(vals[3]).strip() if vals[3] is not None else ''

        if name and sid:
            if not current_col:
                current_col = 'DYPIMED'

            records_by_college[current_col].append({
                'row': r,
                'name': name.upper(),
                'division': div,
                'id': sid,
                'barcode': sid.replace(' ', '').upper(),
                'college': current_col
            })

    return records_by_college

def create_pass_image(template_img, name, student_id, font_path, code128_class, writer):
    """
    Renders a single high-resolution Dandiya pass (3000x900).
    Centers Student Name, Student ID, and Code-128 barcode cleanly on the stub.
    """
    pass_img = template_img.copy()
    draw = ImageDraw.Draw(pass_img)

    center_x = 2419
    max_name_width = 800

    # 1. Dynamic font sizing for student name to avoid overflow
    font_size_name = 100
    font_name = ImageFont.truetype(font_path, font_size_name)
    bb = draw.textbbox((0, 0), name, font=font_name, stroke_width=2)
    while (bb[2] - bb[0]) > max_name_width and font_size_name > 30:
        font_size_name -= 3
        font_name = ImageFont.truetype(font_path, font_size_name)
        bb = draw.textbbox((0, 0), name, font=font_name, stroke_width=2)

    # 2. Font for Student ID / Roll Number
    font_id = ImageFont.truetype(font_path, 120)

    # Position calculations
    bb_n = draw.textbbox((0, 0), name, font=font_name, stroke_width=2)
    wn, hn = bb_n[2] - bb_n[0], bb_n[3] - bb_n[1]
    xn = center_x - wn // 2
    yn = 405 - hn // 2

    bb_i = draw.textbbox((0, 0), student_id, font=font_id, stroke_width=2)
    wi, hi = bb_i[2] - bb_i[0], bb_i[3] - bb_i[1]
    xi = center_x - wi // 2
    yi = 560 - hi // 2

    # Draw Name & ID with crisp white stroke for maximum contrast
    draw.text((xn, yn), name, fill=(255, 255, 255), stroke_width=2, stroke_fill=(255, 255, 255), font=font_name)
    draw.text((xi, yi), student_id, fill=(255, 255, 255), stroke_width=2, stroke_fill=(255, 255, 255), font=font_id)

    # 3. Generate Code 128 barcode
    bc_opts = {
        'module_width': 0.40,
        'module_height': 14.0,
        'quiet_zone': 2.0,
        'font_size': 0,
        'text_distance': 0,
        'write_text': False
    }
    bc = code128_class(student_id, writer=writer)
    fp = io.BytesIO()
    bc.write(fp, options=bc_opts)
    fp.seek(0)
    bc_img = Image.open(fp)

    # Bounding box for barcode within template's white area
    box_w = 580
    box_h = 135
    box_center_x = 2419
    box_center_y = 754

    bw, bh = bc_img.size
    scale = min(box_w / bw, box_h / bh)
    new_bw = max(1, int(bw * scale))
    new_bh = max(1, int(bh * scale))
    bc_resized = bc_img.resize((new_bw, new_bh), Image.Resampling.LANCZOS)

    paste_x = box_center_x - new_bw // 2
    paste_y = box_center_y - new_bh // 2
    pass_img.paste(bc_resized, (paste_x, paste_y))

    return pass_img

def generate_college_passes(college, students, font_path, code128_class, writer):
    """
    Generates both A4 Printable PDF and Multi-Page Strip PDF for a specific college.
    """
    template_path = TEMPLATES[college]
    safe_col = college.replace(' ', '_')

    print(f"\n==========================================")
    print(f"[{college}] Processing {len(students)} passes...")
    print(f"Template: {os.path.basename(template_path)}")
    print(f"==========================================")

    template_img = Image.open(template_path).convert('RGB')

    # Pre-render pass images
    pass_images = []
    for idx, s in enumerate(students):
        pass_img = create_pass_image(template_img, s['name'], s['id'], font_path, code128_class, writer)
        buf = io.BytesIO()
        pass_img.save(buf, format='JPEG', quality=90)
        pass_images.append(buf.getvalue())
        if (idx + 1) % 100 == 0 or (idx + 1) == len(students):
            print(f"  Rendered {idx + 1}/{len(students)} passes...")

    # 1. Printable A4 PDF (3 passes per page with cutting guide borders)
    pdf_a4_name = f"{safe_col}_Passes_A4_Printable.pdf"
    pdf_a4_path = os.path.join(DANDIYAA_DIR, pdf_a4_name)
    print(f"[{college}] Building A4 Printable PDF -> {pdf_a4_path}")

    doc_a4 = fitz.open()
    a4_w, a4_h = 595.28, 841.89  # A4 Portrait in pt
    pass_w = 555.0
    pass_h = pass_w * (900.0 / 3000.0)  # 166.5 pt
    margin_x = (a4_w - pass_w) / 2.0
    spacing_y = 60.0
    start_y = (a4_h - (3 * pass_h + 2 * spacing_y)) / 2.0

    current_page = None
    for i, img_bytes in enumerate(pass_images):
        slot_idx = i % 3
        if slot_idx == 0:
            current_page = doc_a4.new_page(width=a4_w, height=a4_h)

        slot_y = start_y + slot_idx * (pass_h + spacing_y)
        rect_slot = fitz.Rect(margin_x, slot_y, margin_x + pass_w, slot_y + pass_h)
        current_page.insert_image(rect_slot, stream=img_bytes)

        # Light cutting guideline around pass
        border_rect = fitz.Rect(margin_x - 0.5, slot_y - 0.5, margin_x + pass_w + 0.5, slot_y + pass_h + 0.5)
        current_page.draw_rect(border_rect, color=(0.8, 0.8, 0.8), width=0.5)

    doc_a4.save(pdf_a4_path, deflate=True)
    doc_a4.close()
    total_a4_pages = (len(students) + 2) // 3
    a4_size_mb = os.path.getsize(pdf_a4_path) / (1024 * 1024)
    print(f"[{college}] Finished A4 PDF: {total_a4_pages} pages ({a4_size_mb:.2f} MB)")

    # 2. Multi-Page Strip PDF (10" x 3", 720x216 pt)
    pdf_strip_name = f"{safe_col}_Passes.pdf"
    pdf_strip_path = os.path.join(DANDIYAA_DIR, pdf_strip_name)
    print(f"[{college}] Building Strip PDF -> {pdf_strip_path}")

    doc_single = fitz.open()
    strip_w, strip_h = 720, 216
    rect_single = fitz.Rect(0, 0, strip_w, strip_h)

    for img_bytes in pass_images:
        page = doc_single.new_page(width=strip_w, height=strip_h)
        page.insert_image(rect_single, stream=img_bytes)

    doc_single.save(pdf_strip_path, deflate=True)
    doc_single.close()
    strip_size_mb = os.path.getsize(pdf_strip_path) / (1024 * 1024)
    print(f"[{college}] Finished Strip PDF: {len(students)} pages ({strip_size_mb:.2f} MB)")

    # Copy to passes/dandiyaa directory for centralized archive
    os.makedirs(PASSES_COPY_DIR, exist_ok=True)
    shutil.copy2(pdf_a4_path, os.path.join(PASSES_COPY_DIR, pdf_a4_name))
    shutil.copy2(pdf_strip_path, os.path.join(PASSES_COPY_DIR, pdf_strip_name))

    return {
        'college': college,
        'count': len(students),
        'a4_pdf': pdf_a4_path,
        'a4_pages': total_a4_pages,
        'a4_size_mb': a4_size_mb,
        'strip_pdf': pdf_strip_path,
        'strip_size_mb': strip_size_mb
    }

def sync_database(records_by_college):
    """
    Syncs the Dandiya students to database.json so the live scanner web app
    recognizes all passes upon scanning.
    Preserves any existing status (ENTERED, volunteer, timestamps).
    """
    db_path = os.path.join(BASE_DIR, 'database.json')
    db = TinyDB(db_path)
    students_table = db.table('students')
    existing_students = students_table.all()

    existing_map = {}
    for s in existing_students:
        c = s.get('college')
        bc = s.get('barcode', '').strip().upper()
        if c and bc:
            existing_map[(c, bc)] = s

    q = Query()
    new_inserts = []
    updated_count = 0

    for col, recs in records_by_college.items():
        for r in recs:
            key = (col, r['barcode'])
            if key in existing_map:
                ex = existing_map[key]
                if len(r['name']) > len(ex.get('name', '')):
                    students_table.update({'name': r['name']}, (q.college == col) & (q.barcode == r['barcode']))
                    updated_count += 1
            else:
                new_record = {
                    'barcode': r['id'].strip().upper(),
                    'name': r['name'].strip().upper(),
                    'college': col,
                    'division': r['division'],
                    'roll_no': r['id'].strip(),
                    'status': 'PENDING',
                    'entry_timestamp': None,
                    'entry_date': None,
                    'entry_time': None,
                    'volunteer': None
                }
                new_inserts.append(new_record)
                existing_map[key] = new_record

    if new_inserts:
        students_table.insert_multiple(new_inserts)

    db.close()
    print(f"\n[Database Sync] Added {len(new_inserts)} new Dandiya students, updated {updated_count} names.")

def main():
    print("=" * 60)
    print("*** DANDIYA & GARBA NIGHT 2026 - PASS GENERATION SUITE ***")
    print("=" * 60)

    font_path = get_font_path()
    code128_class = barcode.get_barcode_class('code128')
    writer = ImageWriter()

    print(f"Font loaded: {font_path}")
    print(f"Reading Sheet: {SHEET_PATH}")

    records_by_college = load_sheet_records()
    total_records = sum(len(v) for v in records_by_college.values())
    print(f"Total students across all 5 colleges: {total_records}")
    for col, recs in records_by_college.items():
        print(f"  - {col}: {len(recs)} students")

    results = []
    for col in ['DYPIMED', 'DYPCOEI', 'DYPTC', 'DYPSOA', 'BOX OFFICE']:
        students = records_by_college[col]
        res = generate_college_passes(col, students, font_path, code128_class, writer)
        results.append(res)

    # Sync database with Dandiya attendees
    sync_database(records_by_college)

    print("\n" + "=" * 60)
    print("*** ALL 5 COLLEGE DANDIYA PASSES GENERATED SUCCESSFULLY! ***")
    print("=" * 60)
    for r in results:
        print(f"{r['college']:<12}: {r['count']} passes | A4: {r['a4_pages']} pages ({r['a4_size_mb']:.1f} MB) | Strip: ({r['strip_size_mb']:.1f} MB)")
        print(f"   A4 PDF   : {r['a4_pdf']}")
        print(f"   Strip PDF: {r['strip_pdf']}")

if __name__ == '__main__':
    main()
