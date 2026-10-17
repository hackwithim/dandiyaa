import os
import shutil
import re
import openpyxl
from tinydb import TinyDB, Query

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

if os.environ.get('VERCEL'):
    DB_PATH = '/tmp/database.json'
    bundled_db = os.path.join(BASE_DIR, 'database.json')
    if not os.path.exists(DB_PATH) and os.path.exists(bundled_db):
        shutil.copy2(bundled_db, DB_PATH)
else:
    DB_PATH = os.path.join(BASE_DIR, 'database.json')

SHEET_PATH = os.path.join(BASE_DIR, 'Sheet.xlsx')

ADMIN_EMAIL = 'rccindia@webworksstudios.com'
ADMIN_PASS = 'DYPWEBP2026'

def get_db():
    return TinyDB(DB_PATH)

def init_database(force_reinit=False):
    """
    Initializes database.json specifically for the Dandiya & Garba Night 2026 event.
    Seeds the admin user with rccindia@webworksstudios.com / DYPWEBP2026,
    default volunteer gate stations, and all 1,362 attendees from Sheet.xlsx.
    """
    if os.path.exists(DB_PATH) and not force_reinit:
        # Check if users and students exist
        db = TinyDB(DB_PATH)
        users = db.table('users').all()
        students = db.table('students').all()
        db.close()
        if len(users) > 0 and len(students) > 0:
            print(f"[db_loader] Database already initialized ({len(students)} attendees, {len(users)} users).")
            return

    db = TinyDB(DB_PATH)
    users_table = db.table('users')
    students_table = db.table('students')
    logs_table = db.table('logs')

    # 1. Seed Users
    users_table.truncate()
    users_table.insert_multiple([
        {
            'username': ADMIN_EMAIL,
            'password': ADMIN_PASS,
            'role': 'admin',
            'name': 'Chief Event Administrator',
            'created_at': '2026-10-13T20:00:00'
        },
        {
            'username': 'gate1',
            'password': 'dandiya2026',
            'role': 'volunteer',
            'name': 'Gate 1 (Main Arena Entrance)',
            'created_at': '2026-10-13T20:00:00'
        },
        {
            'username': 'gate2',
            'password': 'dandiya2026',
            'role': 'volunteer',
            'name': 'Gate 2 (Fast-Track & VIP)',
            'created_at': '2026-10-13T20:00:00'
        },
        {
            'username': 'gate3',
            'password': 'dandiya2026',
            'role': 'volunteer',
            'name': 'Gate 3 (Box Office Registration)',
            'created_at': '2026-10-13T20:00:00'
        }
    ])

    # 2. Seed Attendees from Sheet.xlsx
    wb = openpyxl.load_workbook(SHEET_PATH, data_only=True)
    ws = wb['Sheet1']

    students_table.truncate()
    students_list = []
    current_col = None

    for r in range(1, ws.max_row + 1):
        vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        c0 = str(vals[0]).strip() if vals[0] is not None else ''

        for known in ['DYPIMED', 'DYPCOEI', 'DYPTC', 'DYPSOA', 'BOX OFFICE']:
            if c0 == known or (known in c0 and len([x for x in vals if x is not None]) <= 2):
                current_col = known
                break

        name = str(vals[1]).strip() if vals[1] is not None else ''
        div = str(vals[2]).strip() if vals[2] is not None else ''
        sid = str(vals[3]).strip() if vals[3] is not None else ''

        if name and sid:
            if not current_col:
                current_col = 'DYPIMED'

            students_list.append({
                'barcode': sid.strip().upper(),
                'name': name.strip().upper(),
                'college': current_col,
                'division': div.strip(),
                'roll_no': sid.strip(),
                'status': 'PENDING',
                'entry_timestamp': None,
                'entry_date': None,
                'entry_time': None,
                'volunteer': None
            })

    students_table.insert_multiple(students_list)
    logs_table.truncate()
    db.close()

    print(f"[db_loader] Successfully initialized {len(students_list)} attendees across 5 colleges.")
    print(f"[db_loader] Admin: {ADMIN_EMAIL} configured successfully.")

if __name__ == '__main__':
    init_database(force_reinit=True)
